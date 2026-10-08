"""Build Excel upload templates with the same choices as the entry forms."""

from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from utils.food_catalog import FOOD_ITEMS, units_for_item, categories_for_item_and_unit


TEMPLATE_ROWS = 5000


def _add_food_cascade_options(workbook, options_sheet, headers):
    """Add named ranges and row-aware formulas for Food Database dropdowns."""
    item_col = get_column_letter(headers.index("Food_Item") + 1)
    unit_col = get_column_letter(headers.index("Purchase_Unit") + 1)
    option_col = 20

    for row, item in enumerate(FOOD_ITEMS, start=1):
        options_sheet.cell(row=row, column=option_col, value=item)
    workbook.defined_names.add(DefinedName(
        "FoodItems",
        attr_text=f"'Options'!${get_column_letter(option_col)}$1:${get_column_letter(option_col)}${len(FOOD_ITEMS)}",
    ))

    next_col = option_col + 1
    for item_index, item in enumerate(FOOD_ITEMS, start=1):
        units = units_for_item(item)
        unit_column = next_col
        next_col += 1
        unit_letter = get_column_letter(unit_column)
        for row, unit in enumerate(units, start=1):
            options_sheet.cell(row=row, column=unit_column, value=unit)
        unit_name = f"DV_UNIT_{item_index}"
        workbook.defined_names.add(DefinedName(
            unit_name,
            attr_text=f"'Options'!${unit_letter}$1:${unit_letter}${len(units)}",
        ))
        for unit_index, unit in enumerate(units, start=1):
            categories = categories_for_item_and_unit(item, unit)
            category_column = next_col
            next_col += 1
            category_letter = get_column_letter(category_column)
            for row, category in enumerate(categories, start=1):
                options_sheet.cell(row=row, column=category_column, value=category)
            category_name = f"DV_CAT_{item_index}_{unit_index}"
            workbook.defined_names.add(DefinedName(
                category_name,
                attr_text=f"'Options'!${category_letter}$1:${category_letter}${len(categories)}",
            ))
    return {
        "food_item": "=FoodItems",
        "purchase_unit": f'=INDIRECT("DV_UNIT_"&MATCH(${item_col}2,FoodItems,0))',
        "food_category": (
            f'=INDIRECT("DV_CAT_"&MATCH(${item_col}2,FoodItems,0)&"_"&'
            f'MATCH(${unit_col}2,INDIRECT("DV_UNIT_"&MATCH(${item_col}2,FoodItems,0)),0))'
        ),
    }


def build_excel_template(dataset_name, schema):
    """Return a formatted .xlsx template for a dataset schema."""
    fields = [field for field in schema["fields"] if field["type"] != "computed"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Template"
    options_sheet = workbook.create_sheet("Options")
    options_sheet.sheet_state = "hidden"

    headers = [field["name"] for field in fields]
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    sheet.row_dimensions[1].height = 24

    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    cascade_formulas = (
        _add_food_cascade_options(workbook, options_sheet, headers)
        if dataset_name == "Food Database"
        else {}
    )

    for column, field in enumerate(fields, start=1):
        letter = get_column_letter(column)
        sheet.column_dimensions[letter].width = min(max(len(field["name"]) + 3, 16), 30)
        validation = None
        if field.get("cascade") in cascade_formulas:
            validation = DataValidation(
                type="list",
                formula1=cascade_formulas[field["cascade"]],
                allow_blank=not field.get("required", False),
            )
            validation.errorTitle = "Invalid selection"
            validation.error = "Choose a value from the dropdown for the selected food item and unit."
            validation.promptTitle = field["name"].replace("_", " ")
            validation.prompt = "This dropdown depends on the previous selection."
        elif field["type"] == "select":
            list_column = column
            for row, option in enumerate(field["options"], start=1):
                options_sheet.cell(row=row, column=list_column, value=option)
            list_name = "DV_" + re.sub(r"[^A-Za-z0-9_]", "_", field["name"])
            options_range = f"'Options'!${letter}$1:${letter}${len(field['options'])}"
            workbook.defined_names.add(DefinedName(list_name, attr_text=options_range))
            validation = DataValidation(type="list", formula1=f"={list_name}", allow_blank=not field.get("required", False))
            validation.errorTitle = "Invalid choice"
            validation.error = "Choose a value from the dropdown list."
            validation.promptTitle = field["name"].replace("_", " ")
            validation.prompt = "Select an option from the list."
        elif field["type"] == "date":
            validation = DataValidation(
                type="date", operator="between", formula1="DATE(1900,1,1)",
                formula2="DATE(9999,12,31)", allow_blank=not field.get("required", False),
            )
            validation.errorTitle = "Invalid date"
            validation.error = "Enter a valid Excel date."
            validation.promptTitle = "Date"
            validation.prompt = "Enter a date; display format is YYYY-MM-DD."
            sheet.column_dimensions[letter].width = 16
            for row in range(2, TEMPLATE_ROWS + 1):
                sheet.cell(row=row, column=column).number_format = "yyyy-mm-dd"

        if validation is not None:
            validation.showErrorMessage = True
            validation.showInputMessage = True
            sheet.add_data_validation(validation)
            validation.add(f"{letter}2:{letter}{TEMPLATE_ROWS}")

    result = BytesIO()
    workbook.save(result)
    result.seek(0)
    return result.getvalue()
